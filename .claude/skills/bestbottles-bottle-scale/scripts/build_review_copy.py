"""Review copy of the bottle measurements for someone to check and correct (sent to Jordan's boss 2026-09-16).

    python3 build_review_copy.py [snapshot dir]
    -> <snapshot dir>/Best Bottles bottle heights - please check.xlsx  and  review-meta.json

Tabs: Start here (steps, definitions, example, progress) · Issues to fix (red; every flagged glass as one plain
question naming its SKUs) · All bottles (Correct / Wrong / Not sure for the rest; flagged rows shaded red and
pointed to the Issues tab so each glass has ONE place to answer) · SKU detail (read only).
Bottle # = the body id in bodies.json; reconcile_review.py reads answers back by that number.
"""
import collections, json, math, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.properties import CalcProperties
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bottle_bodies import build, mm, body_key, variant, default_snapshot

SNAP = sys.argv[1] if len(sys.argv) > 1 else default_snapshot()

d, rows, bodies = build(SNAP)
by_key = collections.defaultdict(list)
for p in rows:
    by_key[body_key(p)].append(p)
key_of = lambda b: (b["family"], b["capacityMl"], b["neck"], b["glass"], b["shape"])

F = "Arial"
def font(**k): return Font(name=F, size=k.pop("size", 10), **k)
HEAD = font(bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="2B2F2D"); RED_HEAD = PatternFill("solid", fgColor="9C2F1E")
INPUT = PatternFill("solid", fgColor="FFF2B3"); INPUT_HEAD = PatternFill("solid", fgColor="8A6D00")
ISSUE = PatternFill("solid", fgColor="F8D9D2"); GREY = PatternFill("solid", fgColor="EFEDE8")
LINE = Side(style="thin", color="D9D3C7"); WRAP = Alignment(wrap_text=True, vertical="top"); TOP = Alignment(vertical="top")

def name(b):
    bits = [b["family"], f'{b["capacityMl"]:g} ml' if b["capacityMl"] else "no capacity", b["neck"] or "no neck finish", b["glass"]]
    if b["shape"]: bits.append(b["shape"])
    return " · ".join(bits)

def odd_skus(b):
    truth = b["heightNoCapMm"]
    skus = by_key[key_of(b)]
    odd = collections.defaultdict(list); blank = []
    for p in skus:
        v = mm(p["heightWithoutCap"])[0]
        if v is None: blank.append(p["websiteSku"])
        elif v != truth: odd[v].append(p["websiteSku"])
    return odd, blank, len(skus)

def question(b):
    odd, blank, n = odd_skus(b)
    truth = b["heightNoCapMm"]
    if b["status"].startswith("Check"):
        if "GBSpry1ozGl" in b["governingSku"] or "slim 30 ml tube" in b["notes"]:
            return ("Height can't be right", f"Convex says {truth:g} mm tall and {b['diameterMm']:g} mm wide, but this is a slim 30 ml spray tube. What are its real height and diameter?")
        return ("Height can't be right", f"Convex says {truth:g} mm tall and {b['diameterMm']:g} mm wide. A glass that size holds about "
                f"{math.pi * (b['diameterMm'] / 2) ** 2 * truth / 1000 * 0.6:.0f} ml, not {b['capacityMl']:g} ml. Is the height or the capacity wrong?")
    if b["status"] == "Heights disagree":
        skus = by_key[key_of(b)]
        majority = [p["websiteSku"] for p in skus if mm(p["heightWithoutCap"])[0] == truth]
        odd_all = [s for ss in odd.values() for s in ss]
        if len(majority) == 1 and len(odd_all) == 1:
            (v, ss), = odd.items()
            return ("Two heights on one bottle", f"The two SKUs disagree: {majority[0]} says {truth:g} mm, {ss[0]} says {v:g} mm. Which is right?")
        parts = [f"{len(majority)} SKU{'s say' if len(majority) != 1 else ' says'} {truth:g} mm"]
        for v, s in sorted(odd.items()):
            parts.append(f"{len(s)} say{'s' if len(s) == 1 else ''} {v:g} mm")
        many = len(odd_all) > 1
        if all(s.endswith("Sht") for s in odd_all):
            tail = (" All the odd ones are short-cap listings: were they measured with the cap on?" if many
                    else " The odd one is a short-cap listing: was it measured with the cap on?")
        else:
            tail = " Are the odd ones typos, or a different bottle?" if many else " Is the odd one a typo, or a different bottle?"
        return ("Two heights on one bottle", "; ".join(parts) + "." + tail)
    if b["status"] == "Some SKUs blank":
        return ("Height missing", f"{len(blank)} of {n} SKUs have no height recorded; the rest say {truth:g} mm. Is {truth:g} mm right for all of them?")
    return ("", "")

def listed(b):
    odd, blank, _ = odd_skus(b)
    skus = by_key[key_of(b)]
    if b["status"].startswith("Check"):
        return "all: " + ", ".join(sorted(p["websiteSku"] for p in skus))
    out = []
    majority = sorted(p["websiteSku"] for p in skus if mm(p["heightWithoutCap"])[0] == b["heightNoCapMm"])
    if majority and len(majority) <= 3:
        out.append(f'{b["heightNoCapMm"]:g} mm: ' + ", ".join(majority))
    for v, s in sorted(odd.items()):
        out.append(f"{v:g} mm: " + ", ".join(sorted(s)))
    if blank:
        out.append(f"blank ({len(blank)}): " + ", ".join(sorted(blank)[:12]) + (f", +{len(blank) - 12} more" if len(blank) > 12 else ""))
    return "\n".join(out)

flagged = [b for b in bodies if b["status"] != "OK"]
order = {"Height can't be right": 0, "Two heights on one bottle": 1, "Height missing": 2}
flagged.sort(key=lambda b: (order[question(b)[0]], -b["skuCount"], b["family"]))

wb = Workbook()
# ============================================================ Start here
st = wb.active; st.title = "Start here"; st.sheet_properties.tabColor = "2B2F2D"
st.column_dimensions["A"].width = 4; st.column_dimensions["B"].width = 34; st.column_dimensions["C"].width = 18; st.column_dimensions["D"].width = 90
st["B2"] = "Bottle heights: please check"; st["B2"].font = font(size=18, bold=True)
st["B3"] = f"These are the bottle measurements the new website uses to size every product photo. Taken from our product database (Convex) on {d['pulledAt'][:10]}."
st["B3"].font = font(color="6F7571")
r = 5
st.cell(row=r, column=2, value="What to do").font = font(size=13, bold=True); r += 1
steps = [
    ("1  Issues to fix (red tab)", "19 bottles where our numbers disagree or can't be right. Each row asks one question and names the SKUs. Please answer these first."),
    ("2  All bottles", "The other 89 bottles. Pick Correct / Wrong / Not sure for each one; if Wrong, type the right number. Start with the rows marked High volume."),
    ("3  Send the file back", "Reply to the email with this file attached. Please don't rename, delete or re-order rows: the Bottle # is how your answers are loaded back."),
]
for t, s in steps:
    st.cell(row=r, column=2, value=t).font = font(bold=True); c = st.cell(row=r, column=4, value=s); c.font = font(); c.alignment = WRAP; r += 1
r += 1
st.cell(row=r, column=2, value="How to fill it in").font = font(size=13, bold=True); r += 1
legend = st.cell(row=r, column=2, value="Yellow cells"); legend.fill = INPUT; legend.font = font(bold=True)
st.cell(row=r, column=4, value="The only cells to type in. Pick from the dropdown where there is one. Measurements in millimetres, numbers only (e.g. 70 or 50.8).").font = font(); r += 1
red = st.cell(row=r, column=2, value="Red rows"); red.fill = ISSUE; red.font = font(bold=True)
st.cell(row=r, column=4, value="On the All bottles tab, a red row is one of the 19 issues: answer it on the Issues to fix tab instead.").font = font(); r += 1
r += 1
st.cell(row=r, column=2, value="What we measure").font = font(size=13, bold=True); r += 1
for t, s in [("Height without cap", "Bare glass, from the bottom of the bottle to the top of the neck. Nothing fitted: no cap, sprayer, pump or roller."),
             ("Height with cap", "The bottle with its cap or fitment on, bottom to the very top."),
             ("Diameter", "The widest part of the bottle."),
             ("One bottle, many SKUs", "Every SKU built on the same glass (clear, amber, cobalt; any cap or fitment) must have the same height without cap.")]:
    st.cell(row=r, column=2, value=t).font = font(bold=True); c = st.cell(row=r, column=4, value=s); c.font = font(); c.alignment = WRAP; r += 1
r += 1
st.cell(row=r, column=2, value="Example of a filled-in issue (not real data)").font = font(size=13, bold=True); r += 1
ex_heads = ["Right height without cap (mm)", "Right height with cap (mm)", "Right diameter (mm)", "The odd SKUs are", "How you checked", "Comment"]
ex_vals = [53, 72, 17, "A typo, same bottle", "Measured a sample", "Same bottle as the rest. The 60 mm figure was measured with the cap on."]
for k, (h, v) in enumerate(zip(ex_heads, ex_vals)):
    st.cell(row=r + k, column=2, value=h).font = font(color="6F7571")
    c = st.cell(row=r + k, column=4, value=v); c.fill = INPUT; c.font = font(italic=True); c.alignment = Alignment(horizontal="left")
r += len(ex_heads) + 1
st.cell(row=r, column=2, value="Progress").font = font(size=13, bold=True); prog_row = r + 1

# ============================================================ Issues to fix
iss = wb.create_sheet("Issues to fix"); iss.sheet_properties.tabColor = "C0392B"
icols = [("Bottle #", 8, None), ("Issue", 20, None), ("Bottle", 34, None), ("Question for you", 58, None), ("SKUs involved", 44, None),
         ("Governing SKU", 22, None), ("Height without cap now (mm)", 13, None), ("Height with cap now (mm)", 13, None), ("Diameter now (mm)", 11, None),
         ("SKUs on this bottle", 10, None),
         ("Right height without cap (mm)", 14, "in"), ("Right height with cap (mm)", 13, "in"), ("Right diameter (mm)", 12, "in"), ("Right capacity (ml)", 11, "in"),
         ("The odd SKUs are", 20, "in"), ("How you checked", 18, "in"), ("Comment", 40, "in"), ("Answered?", 16, "calc")]
for c, (h, w, kind) in enumerate(icols, 1):
    cell = iss.cell(row=1, column=c, value=h); cell.font = HEAD; cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.fill = INPUT_HEAD if kind == "in" else (RED_HEAD if kind is None else HFILL); iss.column_dimensions[L(c)].width = w
iss.row_dimensions[1].height = 44
for i, b in enumerate(flagged, 2):
    kind, q = question(b)
    vals = [b["id"], kind, name(b), q, listed(b) or "—", b["governingSku"], b["heightNoCapMm"],
            f'{b["heightCapMinMm"]:g}–{b["heightCapMaxMm"]:g}' if b["heightCapMinMm"] is not None and b["heightCapMinMm"] != b["heightCapMaxMm"] else b["heightCapMinMm"],
            b["diameterMm"], b["skuCount"]]
    for c, v in enumerate(vals, 1):
        cell = iss.cell(row=i, column=c, value=v); cell.font = font(bold=(c in (2, 4))); cell.alignment = WRAP; cell.border = Border(bottom=LINE)
    for c in range(11, 18):
        cell = iss.cell(row=i, column=c); cell.fill = INPUT; cell.border = Border(bottom=LINE, left=LINE); cell.alignment = WRAP
    iss.cell(row=i, column=18, value=f'=IF(COUNTA(K{i}:Q{i})=0,"To answer",IF(AND(K{i}="",Q{i}=""),"Add the right height","Answered"))').font = font(bold=True)
    iss.row_dimensions[i].height = max(48, 15 * (1 + len(q) // 60 + listed(b).count("\n")))
last_i = len(flagged) + 1
iss.freeze_panes = "D2"; iss.auto_filter.ref = f"A1:R{last_i}"
dv_odd = DataValidation(type="list", formula1='"A typo, same bottle,A different bottle,Not sure,Not applicable"', allow_blank=True)
dv_how = DataValidation(type="list", formula1='"Spec sheet,Measured a sample,Asked the supplier,From memory,Other"', allow_blank=True)
dv_mm = DataValidation(type="decimal", operator="between", formula1="1", formula2="400", allow_blank=True, error="Millimetres, numbers only", errorTitle="Measurement")
dv_ml = DataValidation(type="decimal", operator="between", formula1="0.5", formula2="2000", allow_blank=True, error="Millilitres, numbers only", errorTitle="Capacity")
for dv in (dv_odd, dv_how, dv_mm, dv_ml): iss.add_data_validation(dv)
dv_odd.add(f"O2:O{last_i}"); dv_how.add(f"P2:P{last_i}"); dv_mm.add(f"K2:M{last_i}"); dv_ml.add(f"N2:N{last_i}")
iss.conditional_formatting.add(f"R2:R{last_i}", FormulaRule(formula=['$R2="Answered"'], font=Font(name=F, size=10, bold=True, color="2E7D50")))
iss.conditional_formatting.add(f"R2:R{last_i}", FormulaRule(formula=['$R2<>"Answered"'], font=Font(name=F, size=10, bold=True, color="9C2F1E")))

# ============================================================ All bottles
al = wb.create_sheet("All bottles"); al.sheet_properties.tabColor = "8A6D00"
acols = [("Bottle #", 8, None), ("Priority", 14, None), ("Bottle", 40, None), ("Governing SKU", 24, None), ("Product", 44, None), ("SKUs on this bottle", 10, None),
         ("Height without cap now (mm)", 13, None), ("Height with cap now, min (mm)", 13, None), ("Height with cap now, max (mm)", 13, None), ("Diameter now (mm)", 11, None),
         ("Is the height without cap right?", 16, "in"), ("Right height without cap (mm)", 14, "in"), ("Right height with cap (mm)", 13, "in"), ("Right diameter (mm)", 12, "in"),
         ("How you checked", 18, "in"), ("Comment", 36, "in"), ("Review status", 20, "calc")]
for c, (h, w, kind) in enumerate(acols, 1):
    cell = al.cell(row=1, column=c, value=h); cell.font = HEAD; cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.fill = INPUT_HEAD if kind == "in" else HFILL; al.column_dimensions[L(c)].width = w
al.row_dimensions[1].height = 44
flag_ids = {b["id"] for b in flagged}
def prio(b):
    if b["id"] in flag_ids: return "Issue"
    return "High volume" if b["skuCount"] >= 30 else "Spot check"
ordered = sorted(bodies, key=lambda b: ({"Issue": 0, "High volume": 1, "Spot check": 2}[prio(b)], b["family"], b["capacityMl"] or 0, b["neck"], b["glass"]))
for i, b in enumerate(ordered, 2):
    vals = [b["id"], prio(b), name(b), b["governingSku"], b["governingItem"], b["skuCount"], b["heightNoCapMm"], b["heightCapMinMm"], b["heightCapMaxMm"], b["diameterMm"]]
    for c, v in enumerate(vals, 1):
        cell = al.cell(row=i, column=c, value=v); cell.font = font(bold=(c == 7)); cell.alignment = WRAP if c in (3, 5) else TOP; cell.border = Border(bottom=LINE)
        if b["id"] in flag_ids: cell.fill = ISSUE
    if b["id"] in flag_ids:
        al.cell(row=i, column=11, value="See Issues to fix tab").font = font(bold=True, color="9C2F1E")
        for c in range(11, 17): al.cell(row=i, column=c).fill = ISSUE
        al.cell(row=i, column=17, value="On the Issues tab").font = font(color="9C2F1E")
    else:
        for c in range(11, 17):
            cell = al.cell(row=i, column=c); cell.fill = INPUT; cell.border = Border(bottom=LINE, left=LINE); cell.alignment = WRAP
        al.cell(row=i, column=17, value=f'=IF(K{i}="","To check",IF(K{i}="Correct","Confirmed",IF(K{i}="Not sure","Not sure",IF(L{i}="","Add the right height","Correction given"))))').font = font(bold=True)
last_a = len(ordered) + 1
al.freeze_panes = "D2"; al.auto_filter.ref = f"A1:Q{last_a}"
dv_ok = DataValidation(type="list", formula1='"Correct,Wrong,Not sure"', allow_blank=True)
dv_how2 = DataValidation(type="list", formula1='"Spec sheet,Measured a sample,Asked the supplier,From memory,Other"', allow_blank=True)
dv_mm2 = DataValidation(type="decimal", operator="between", formula1="1", formula2="400", allow_blank=True, error="Millimetres, numbers only", errorTitle="Measurement")
for dv in (dv_ok, dv_how2, dv_mm2): al.add_data_validation(dv)
open_rows = [i for i, b in enumerate(ordered, 2) if b["id"] not in flag_ids]
for i in open_rows:
    dv_ok.add(f"K{i}"); dv_how2.add(f"O{i}"); dv_mm2.add(f"L{i}:N{i}")
al.conditional_formatting.add(f"Q2:Q{last_a}", FormulaRule(formula=['OR($Q2="Confirmed",$Q2="Correction given")'], font=Font(name=F, size=10, bold=True, color="2E7D50")))
al.conditional_formatting.add(f"Q2:Q{last_a}", FormulaRule(formula=['OR($Q2="Add the right height",$Q2="Not sure")'], font=Font(name=F, size=10, bold=True, color="9C2F1E")))

# ============================================================ SKU detail (read only)
sk = wb.create_sheet("SKU detail"); sk.sheet_properties.tabColor = "9AA09C"
scols = [("Bottle #", 8), ("Website SKU", 28), ("Product", 58), ("Glass colour", 12), ("Applicator", 22), ("Cap colour", 16),
         ("Height without cap (mm)", 12), ("Height with cap (mm)", 12), ("Diameter (mm)", 10), ("As recorded", 14)]
for c, (h, w) in enumerate(scols, 1):
    cell = sk.cell(row=1, column=c, value=h); cell.font = HEAD; cell.fill = HFILL; cell.alignment = Alignment(wrap_text=True, vertical="center"); sk.column_dimensions[L(c)].width = w
sk.row_dimensions[1].height = 32
id_of = {key_of(b): b["id"] for b in bodies}
srt = sorted(rows, key=lambda p: (id_of[body_key(p)], p["websiteSku"]))
for i, p in enumerate(srt, 2):
    bid = id_of[body_key(p)]
    for c, v in enumerate([bid, p["websiteSku"], p["itemName"], p["color"], p["applicator"], p["capColor"],
                           mm(p["heightWithoutCap"])[0], mm(p["heightWithCap"])[0], mm(p["diameter"])[0], p["heightWithoutCap"]], 1):
        cell = sk.cell(row=i, column=c, value=v); cell.font = font()
        if bid in flag_ids: cell.fill = ISSUE
sk.freeze_panes = "C2"; sk.auto_filter.ref = f"A1:J{len(srt) + 1}"

# ============================================================ progress formulas on Start here
IS, AL = "'Issues to fix'", "'All bottles'"
prog = [("Issues answered", f'=COUNTIF({IS}!R2:R{last_i},"Answered")', f"of {len(flagged)}"),
        ("Bottles confirmed correct", f'=COUNTIF({AL}!Q2:Q{last_a},"Confirmed")', f"of {len(open_rows)}"),
        ("Bottles with a correction", f'=COUNTIF({AL}!Q2:Q{last_a},"Correction given")', ""),
        ("Bottles marked not sure", f'=COUNTIF({AL}!Q2:Q{last_a},"Not sure")', ""),
        ("Bottles still to check", f'=COUNTIF({AL}!Q2:Q{last_a},"To check")', "")]
for k, (label, f, note) in enumerate(prog):
    st.cell(row=prog_row + k, column=2, value=label).font = font()
    c = st.cell(row=prog_row + k, column=3, value=f); c.font = font(bold=True); c.alignment = Alignment(horizontal="right")
    st.cell(row=prog_row + k, column=4, value=note).font = font(color="6F7571")

wb.calculation = CalcProperties(fullCalcOnLoad=True)
OUT = os.path.join(SNAP, "Best Bottles bottle heights - please check.xlsx")
wb.save(OUT)
json.dump(dict(flagged=[b["id"] for b in flagged], lastIssue=last_i, lastAll=last_a, openRows=len(open_rows)), open(os.path.join(SNAP, "review-meta.json"), "w"))
print("saved", OUT, "| issues", len(flagged), "| all bottles", len(ordered), "| to confirm", len(open_rows), "| SKUs", len(srt))
for b in flagged:
    k, q = question(b); print(f"  #{b['id']:>3} {k:26} {name(b):48} {q[:110]}")

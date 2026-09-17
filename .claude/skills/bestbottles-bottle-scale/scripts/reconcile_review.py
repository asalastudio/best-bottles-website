"""Turn a returned review file into proposed Convex corrections. Never writes to Convex.

    python3 reconcile_review.py <returned .xlsx> [snapshot dir]
    -> <snapshot dir>/reconcile-<timestamp>/proposed-corrections.csv   one row per SKU field to change
                                          confirmed.csv                glasses the reviewer marked Correct
                                          follow-up.csv                Not sure, missing numbers, different-bottle calls
                                          summary.md                   counts and what to do next

Reads the file against the SAME snapshot it was built from (bodies.json ids are positional). Rules:
  Issues to fix tab
    * Right height / cap height / diameter given, "The odd SKUs are" = A typo, same bottle (or blank)
        -> every SKU on that glass gets the numbers given.
    * "A different bottle" -> the glass's usual SKUs get the numbers given; the odd SKUs are left alone and
        listed in follow-up (they need their own body and their own measurement).
    * "Not sure", or a row with only a comment -> follow-up, no change.
    * Right capacity (ml) -> proposed capacityMl change for every SKU on the glass (capacity also drives names,
        so treat it as a catalogue change, not just a measurement).
  All bottles tab
    * Correct -> confirmed.  Wrong + numbers -> every SKU on the glass.  Wrong without a number / Not sure -> follow-up.
Proposed values keep the Convex text format "<mm> ±<tolerance> mm" using the glass's existing tolerance.
Applying them needs Jordan's approval and a production write; the canonical measurements live in production.
"""
import collections, csv, json, os, sys
from datetime import datetime
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bottle_bodies import build, default_snapshot, mm, skus_by_body


def fmt(value, tol):
    v = f"{float(value):g}"
    return f"{v} ±{tol:g} mm" if tol else f"{v} mm"


def num(x):
    if x in (None, ""):
        return None
    try:
        return float(str(x).replace("mm", "").replace("ml", "").strip())
    except ValueError:
        return None


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    returned = sys.argv[1]
    snap = sys.argv[2] if len(sys.argv) > 2 else default_snapshot()
    d, rows, bodies = build(snap)
    by_id = {b["id"]: b for b in bodies}
    skus = skus_by_body(rows, bodies)
    wb = load_workbook(returned, data_only=False)

    proposed, confirmed, follow = [], [], []

    def change_all(b, field, new_value, reason, how, comment, only=None):
        tol = b["toleranceMm"]
        for p in skus[b["id"]]:
            if only is not None and p["websiteSku"] not in only:
                continue
            current = p[field]
            if field == "capacityMl":
                new_text = new_value
                if current is not None and float(current) == float(new_value):
                    continue
            else:
                new_text = fmt(new_value, tol)
                if mm(current)[0] == float(new_value):
                    continue
            proposed.append(dict(websiteSku=p["websiteSku"], field=field, current=current, proposed=new_text, bottle=b["id"],
                                 bottleName=f'{b["family"]} {b["capacityMl"]:g} ml {b["neck"]} {b["glass"]} {b["shape"]}'.strip(),
                                 reason=reason, howChecked=how or "", comment=comment or ""))

    # ------------------------------------------------------------ Issues to fix
    ws = wb["Issues to fix"]
    head = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    col = lambda r, name: ws.cell(row=r, column=head[name]).value
    for r in range(2, ws.max_row + 1):
        bid = col(r, "Bottle #")
        if bid is None:
            continue
        b = by_id[int(bid)]
        h, hc, dia, cap = num(col(r, "Right height without cap (mm)")), num(col(r, "Right height with cap (mm)")), \
            num(col(r, "Right diameter (mm)")), num(col(r, "Right capacity (ml)"))
        odd_answer, how, comment = col(r, "The odd SKUs are"), col(r, "How you checked"), col(r, "Comment")
        if not any(v not in (None, "") for v in (h, hc, dia, cap, odd_answer, how, comment)):
            follow.append(dict(bottle=b["id"], websiteSku=b["governingSku"], item="Issue not answered", detail=col(r, "Question for you")))
            continue
        truth = b["heightNoCapMm"]
        odd = {p["websiteSku"] for p in skus[b["id"]] if mm(p["heightWithoutCap"])[0] not in (None, truth)}
        only = None
        if odd_answer == "A different bottle":
            only = {p["websiteSku"] for p in skus[b["id"]]} - odd
            follow.append(dict(bottle=b["id"], websiteSku=", ".join(sorted(odd)), item="Reviewer says these are a different bottle",
                               detail="Give them their own body and confirm their own height; left unchanged"))
        if odd_answer == "Not sure":
            follow.append(dict(bottle=b["id"], websiteSku=b["governingSku"], item="Reviewer not sure", detail=comment or ""))
            continue
        if h is None and hc is None and dia is None and cap is None:
            follow.append(dict(bottle=b["id"], websiteSku=b["governingSku"], item="Answered without a number", detail=f"{odd_answer or ''} {comment or ''}".strip()))
            continue
        reason = f"Issues to fix: {col(r, 'Issue')}"
        if h is not None: change_all(b, "heightWithoutCap", h, reason, how, comment, only)
        if hc is not None: change_all(b, "heightWithCap", hc, reason, how, comment, only)
        if dia is not None: change_all(b, "diameter", dia, reason, how, comment, only)
        if cap is not None: change_all(b, "capacityMl", cap, reason + " (capacity: a catalogue change)", how, comment, only)

    # ------------------------------------------------------------ All bottles
    ws = wb["All bottles"]
    head = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    col = lambda r, name: ws.cell(row=r, column=head[name]).value
    for r in range(2, ws.max_row + 1):
        bid = col(r, "Bottle #")
        if bid is None:
            continue
        b = by_id[int(bid)]
        answer = col(r, "Is the height without cap right?")
        if answer in (None, "", "See Issues to fix tab"):
            continue
        how, comment = col(r, "How you checked"), col(r, "Comment")
        if answer == "Correct":
            confirmed.append(dict(bottle=b["id"], websiteSku=b["governingSku"], heightNoCapMm=b["heightNoCapMm"], howChecked=how or "", comment=comment or ""))
            continue
        if answer == "Not sure":
            follow.append(dict(bottle=b["id"], websiteSku=b["governingSku"], item="Reviewer not sure", detail=comment or ""))
            continue
        h, hc, dia = num(col(r, "Right height without cap (mm)")), num(col(r, "Right height with cap (mm)")), num(col(r, "Right diameter (mm)"))
        if h is None and hc is None and dia is None:
            follow.append(dict(bottle=b["id"], websiteSku=b["governingSku"], item="Marked Wrong without a number", detail=comment or ""))
            continue
        if h is not None: change_all(b, "heightWithoutCap", h, "All bottles: marked Wrong", how, comment)
        if hc is not None: change_all(b, "heightWithCap", hc, "All bottles: marked Wrong", how, comment)
        if dia is not None: change_all(b, "diameter", dia, "All bottles: marked Wrong", how, comment)

    out = os.path.join(snap, f"reconcile-{datetime.now().strftime('%Y%m%d-%H%M%S')}")
    os.makedirs(out, exist_ok=True)
    for name, data, fields in (("proposed-corrections.csv", proposed, ["websiteSku", "field", "current", "proposed", "bottle", "bottleName", "reason", "howChecked", "comment"]),
                               ("confirmed.csv", confirmed, ["bottle", "websiteSku", "heightNoCapMm", "howChecked", "comment"]),
                               ("follow-up.csv", follow, ["bottle", "websiteSku", "item", "detail"])):
        with open(os.path.join(out, name), "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(data)
    fields = collections.Counter(p["field"] for p in proposed)
    touched = len({p["bottle"] for p in proposed})
    summary = [f"# Review reconciliation", "", f"- Returned file: `{os.path.basename(returned)}`",
               f"- Snapshot: `{os.path.relpath(snap)}` (pulled {d['pulledAt']})", "",
               f"- Proposed field changes: **{len(proposed)}** on {touched} glasses ({', '.join(f'{k} {v}' for k, v in fields.items()) or 'none'})",
               f"- Glasses confirmed correct: **{len(confirmed)}**", f"- Follow-ups: **{len(follow)}**", "",
               "Nothing has been written to Convex. Next: Jordan approves proposed-corrections.csv; then write the values to",
               "production (where the canonical measurements live), re-pull the snapshot, and rebuild bodies.json."]
    open(os.path.join(out, "summary.md"), "w").write("\n".join(summary) + "\n")
    print("\n".join(summary))
    print(f"\nwritten to {out}")


if __name__ == "__main__":
    main()

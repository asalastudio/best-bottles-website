"""Measurements workbook from a production snapshot: one row per glass (the source of truth), every SKU, and a read-me.

    python3 build_measurements_workbook.py [snapshot dir]
    -> <snapshot dir>/Best Bottles bottle bodies - measurements (Convex production <date>).xlsx and .csv

No LibreOffice on Jordan's machine, so formulas are not pre-calculated: fullCalcOnLoad is set and Excel or Numbers
calculates them on open. Cross-check any formula you add in Python before shipping.
"""
import csv, json, collections, os, sys
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bottle_bodies import build, mm, glass_kind, is_tall_glass as is_tall, variant, default_snapshot

SNAP = sys.argv[1] if len(sys.argv) > 1 else default_snapshot()

d, rows, bodies = build(SNAP)
pulled = datetime.fromisoformat(d["pulledAt"].replace("Z", "+00:00")).astimezone()
STAMP = pulled.strftime("%Y-%m-%d %H:%M %Z")
NAME = os.path.join(SNAP, f"Best Bottles bottle bodies - measurements (Convex production {d['pulledAt'][:10]})")

F = "Arial"
H_FONT = Font(name=F, bold=True, color="FFFFFF", size=10); H_FILL = PatternFill("solid", fgColor="2B2F2D")
BODY = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True); MUTED = Font(name=F, size=10, color="6F7571")
TITLE = Font(name=F, size=16, bold=True); SUB = Font(name=F, size=10, color="6F7571")
TRUTH_FILL = PatternFill("solid", fgColor="F3EEDF"); THIN = Side(style="thin", color="D9D3C7")
WRAP = Alignment(wrap_text=True, vertical="top"); TOP = Alignment(vertical="top")

wb = Workbook()

# ------------------------------------------------------------------ sheet 1: one row per glass
ws = wb.active; ws.title = "Bottle bodies"
cols = [("#", 5), ("Family", 16), ("Capacity (ml)", 10), ("Neck finish", 12), ("Glass finish", 11), ("Variant", 12),
        ("Height without cap (mm)", 13), ("Tolerance (± mm)", 10), ("Height with cap, min (mm)", 13), ("Height with cap, max (mm)", 13),
        ("Diameter (mm)", 10), ("Governing SKU", 24), ("Governing product", 46), ("SKUs on this glass", 10), ("Glass colours", 22),
        ("Heights recorded across its SKUs", 30), ("SKUs with a different height", 46), ("Measurement source", 26), ("Status", 17), ("Notes", 52)]
for c, (name, w) in enumerate(cols, 1):
    cell = ws.cell(row=1, column=c, value=name); cell.font = H_FONT; cell.fill = H_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="center"); ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[1].height = 42
for i, b in enumerate(bodies, 1):
    r = i + 1
    vals = [i, b["family"], b["capacityMl"], b["neck"] or "—", b["glass"], b["shape"] or "", b["heightNoCapMm"], b["toleranceMm"],
            b["heightCapMinMm"], b["heightCapMaxMm"], b["diameterMm"], b["governingSku"], b["governingItem"], b["skuCount"], b["colours"],
            b["heightsRecorded"], b["oddSkus"], b["sources"], b["status"], b["notes"]]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v); cell.font = BODY; cell.alignment = WRAP if c in (13, 15, 16, 17, 18, 20) else TOP
        cell.border = Border(bottom=THIN)
    ws.cell(row=r, column=7).font = BOLD; ws.cell(row=r, column=7).fill = TRUTH_FILL
    for c in (3, 7, 8, 9, 10, 11):
        ws.cell(row=r, column=c).number_format = "0.0##"
last = len(bodies) + 1
ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
rng = f"S2:S{last}"
ws.conditional_formatting.add(rng, FormulaRule(formula=['$S2="OK"'], font=Font(name=F, size=10, color="2E7D50", bold=True)))
ws.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("Check",$S2))'], fill=PatternFill("solid", fgColor="F6D5CE"), font=Font(name=F, size=10, color="9C2F1E", bold=True)))
ws.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("disagree",$S2))'], fill=PatternFill("solid", fgColor="FBE8C8"), font=Font(name=F, size=10, color="8A5A00", bold=True)))
ws.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("blank",$S2))'], fill=PatternFill("solid", fgColor="E8EEF4"), font=Font(name=F, size=10, color="2B4F72", bold=True)))

# ------------------------------------------------------------------ sheet 2: every SKU, tied to its body
body_no = {}
for i, b in enumerate(bodies, 1):
    body_no[(b["family"], b["capacityMl"], b["neck"], b["glass"], b["shape"])] = i
ws2 = wb.create_sheet("All SKUs")
cols2 = [("Body #", 7), ("Family", 16), ("Capacity (ml)", 10), ("Neck finish", 12), ("Glass finish", 11), ("Glass colour", 12),
         ("Applicator", 22), ("Cap colour", 16), ("Cap height", 9), ("Height without cap (mm)", 13), ("Height with cap (mm)", 12),
         ("Diameter (mm)", 10), ("Height as recorded", 14), ("Measurement source", 26), ("Website SKU", 26), ("Grace SKU", 26), ("Product", 60)]
for c, (name, w) in enumerate(cols2, 1):
    cell = ws2.cell(row=1, column=c, value=name); cell.font = H_FONT; cell.fill = H_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="center"); ws2.column_dimensions[get_column_letter(c)].width = w
ws2.row_dimensions[1].height = 32
shape_of = variant
srt = sorted(rows, key=lambda p: (p["family"] or "", p["capacityMl"] or 0, p["neckThreadSize"] or "", p["websiteSku"]))
for r, p in enumerate(srt, 2):
    key = ((p["family"] or "Unknown"), p["capacityMl"], p["neckThreadSize"] or "", glass_kind(p), variant(p))
    src = p["measurementSource"] or ""
    vals = [body_no.get(key), p["family"], p["capacityMl"], p["neckThreadSize"] or "—", glass_kind(p) + (" · Tall" if is_tall(p) else ""), p["color"],
            p["applicator"], p["capColor"], p["capHeight"], mm(p["heightWithoutCap"])[0], mm(p["heightWithCap"])[0], mm(p["diameter"])[0],
            p["heightWithoutCap"], "master truth 2026-07-12" if src.startswith("best-bottles-master-truth") else ("legacy exact" if src.startswith("legacy-exact") else "none"),
            p["websiteSku"], p["graceSku"], p["itemName"]]
    for c, v in enumerate(vals, 1):
        cell = ws2.cell(row=r, column=c, value=v); cell.font = BODY
    for c in (3, 10, 11, 12):
        ws2.cell(row=r, column=c).number_format = "0.0##"
ws2.freeze_panes = "B2"; ws2.auto_filter.ref = f"A1:{get_column_letter(len(cols2))}{len(srt) + 1}"
missing_body = sum(1 for p in srt if body_no.get(((p["family"] or "Unknown"), p["capacityMl"], p["neckThreadSize"] or "", glass_kind(p), variant(p))) is None)

# ------------------------------------------------------------------ sheet 3: read me + formula counts
ws3 = wb.create_sheet("Read me", 0)
ws3.column_dimensions["A"].width = 38; ws3.column_dimensions["B"].width = 16; ws3.column_dimensions["C"].width = 70
ws3["A1"] = "Best Bottles bottle bodies"; ws3["A1"].font = TITLE
ws3["A2"] = f"Measurements from Convex production (precise-raccoon-123), products table, pulled {STAMP}."; ws3["A2"].font = SUB
ws3["A3"] = "One row per physical glass: the bottle that governs every SKU built on it. Its height without cap is the source of truth for sizing."; ws3["A3"].font = SUB
B = "'Bottle bodies'"; n = last
summary = [
    ("Bottle bodies (glasses)", f"=COUNTA({B}!B2:B{n})", "Family + capacity + neck finish + glass finish (standard, frosted, swirl) + tall glass. Glass colour and fitment do not change the body."),
    ("SKUs covered", f"=SUM({B}!N2:N{n})", "Every live bottle, jar, aluminium, plastic and metal-atomizer SKU. Retired SKUs, caps, components and packaging are excluded."),
    ("Families", "=COUNTA(FAMLIST)", "Listed under Bodies per family below."),
    ("Status OK", f'=COUNTIF({B}!S2:S{n},"OK")', "Every SKU on the glass records the same height without cap."),
    ("Heights disagree", f'=COUNTIF({B}!S2:S{n},"Heights disagree")', "Most common height is used; the odd SKUs are listed so they can be corrected."),
    ("Check height", f'=COUNTIF({B}!S2:S{n},"Check height")', "The recorded height cannot be right for the capacity or the photograph."),
    ("Some SKUs blank", f'=COUNTIF({B}!S2:S{n},"Some SKUs blank")', "The glass has a height, but some of its SKUs have none recorded."),
]
ws3["A5"] = "Summary"; ws3["A5"].font = BOLD
for k, (label, formula, note) in enumerate(summary, 6):
    ws3.cell(row=k, column=1, value=label).font = BODY
    c = ws3.cell(row=k, column=2, value=formula); c.font = BOLD; c.number_format = "0"
    ws3.cell(row=k, column=3, value=note).font = MUTED; ws3.cell(row=k, column=3).alignment = WRAP
row = 6 + len(summary) + 1
ws3.cell(row=row, column=1, value="How to read a row").font = BOLD; row += 1
for label, text in [("Height without cap (mm)", "The source of truth. Bare glass, foot to rim. When SKUs disagree, the value most of them record."),
                    ("Height with cap, min / max (mm)", "Varies with the fitment, so shown as the range across the glass's SKUs."),
                    ("Governing SKU", "The SKU that stands for the glass: records the source-of-truth height, clear glass first, audited measurement first."),
                    ("Measurement source", "master truth = the 2026-07-12 audited master written to production; legacy exact = copied from the legacy site; none = no provenance stamp."),
                    ("Frosted and swirl", "Listed as their own rows because they are different bottles to photograph; several share the clear glass's height."),
                    ("All SKUs sheet", "Every SKU with its body number, so any row can be traced and corrected in Convex.")]:
    ws3.cell(row=row, column=1, value=label).font = BODY
    ws3.cell(row=row, column=3, value=text).font = MUTED; ws3.cell(row=row, column=3).alignment = WRAP; row += 1
row += 1
ws3.cell(row=row, column=1, value="Bodies per family").font = BOLD
fam_first = row + 1
ws3.cell(row=row, column=2, value="Bodies").font = BOLD; ws3.cell(row=row, column=3, value="SKUs").font = BOLD; row += 1
for fam in sorted({b["family"] for b in bodies}):
    ws3.cell(row=row, column=1, value=fam).font = BODY
    ws3.cell(row=row, column=2, value=f"=COUNTIF({B}!$B$2:$B${n},A{row})").font = BODY
    c = ws3.cell(row=row, column=3, value=f"=SUMIF({B}!$B$2:$B${n},A{row},{B}!$N$2:$N${n})"); c.font = BODY; c.alignment = Alignment(horizontal="left")
    row += 1

fam_last = row - 1
for rr in range(6, 6 + len(summary)):
    if ws3.cell(row=rr, column=2).value == "=COUNTA(FAMLIST)":
        ws3.cell(row=rr, column=2, value=f"=COUNTA(A{fam_first}:A{fam_last})")
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
out = f"{NAME}.xlsx"
wb.save(out)
with open(f"{NAME}.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow([c[0] for c in cols])
    for i, b in enumerate(bodies, 1):
        w.writerow([i, b["family"], b["capacityMl"], b["neck"], b["glass"], b["shape"], b["heightNoCapMm"], b["toleranceMm"], b["heightCapMinMm"],
                    b["heightCapMaxMm"], b["diameterMm"], b["governingSku"], b["governingItem"], b["skuCount"], b["colours"], b["heightsRecorded"],
                    b["oddSkus"], b["sources"], b["status"], b["notes"]])
twins = len({(b["family"], b["capacityMl"], b["neck"], b["heightNoCapMm"], b["shape"]) for b in bodies})
print("saved:", out, "| bodies", len(bodies), "| SKUs", len(srt), "| SKUs without a body match", missing_body, "| counting frosted/swirl twins once:", twins)
